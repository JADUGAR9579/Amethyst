"""Authoring and editing the documents that are not text.

The asymmetry with `extract.py` is deliberate. Reading is transparent -- the
model cannot know a path is a PDF before opening it, so `view_file` routes by
extension. Writing is explicit, because the model always knows what it intends
to create, and having `write_file` silently render markdown into Word whenever a
path ended `.docx` would reinterpret content somebody may have meant literally.

Markdown in, document out, and the mapping is the exact inverse of the one
`extract.py` uses, so a document written here and read back gives the markdown
it was written from:

    # Title      the document's Title style / the first sheet / the deck
    ## Heading   Heading 1 / a new sheet / a new slide
    - item       a bullet
    | a | b |    a table

Two traps live in this file, both silent and both total if got wrong. They are
commented where they happen: opening a workbook with `data_only=True` before
saving it, and replacing text run by run in a Word paragraph.
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from pathlib import Path

from backend.documents.extract import LEGACY

#: Formats this module can create from markdown.
WRITABLE = frozenset({".docx", ".xlsx", ".pptx", ".pdf"})

#: Formats whose text can be changed in place. A PDF is not one: it is a fixed
#: layout rather than a document with paragraphs, and PyMuPDF's redact-and-
#: reinsert leaves a visibly broken page while reporting success.
EDITABLE = frozenset({".docx", ".xlsx", ".pptx"})

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")
_TABLE_RULE = re.compile(r"^\s*\|[\s:|-]+\|\s*$")


class DocumentError(RuntimeError):
    """The document could not be written. Carries a sentence fit to show someone."""


@dataclass
class Block:
    kind: str  # heading | para | bullet | table
    text: str = ""
    level: int = 0
    rows: list[list[str]] = field(default_factory=list)


def _unescape(text: str) -> str:
    """Undo what `extract._escape_body` and `extract._cell` did.

    Without this a document read out and written back accumulates a backslash
    before every pipe on each round trip.
    """
    return text.replace("\\|", "|").replace("\\#", "#")


def parse_markdown(text: str) -> list[Block]:
    """Markdown as a flat list of blocks, in document order.

    Deliberately small: headings, bullets, tables and paragraphs, which is the
    set every target format can represent. Anything else -- code fences,
    blockquotes, emphasis -- lands as a paragraph rather than being dropped, so
    content survives even where formatting does not. The tool description says
    so, so nobody discovers it from the output.
    """
    blocks: list[Block] = []
    paragraph: list[str] = []
    table: list[list[str]] = []

    def flush_paragraph() -> None:
        if paragraph:
            blocks.append(Block("para", _unescape(" ".join(paragraph).strip())))
            paragraph.clear()

    def flush_table() -> None:
        if table:
            blocks.append(Block("table", rows=[row[:] for row in table]))
            table.clear()

    for line in text.splitlines():
        if _TABLE_RULE.match(line) and table:
            continue  # the |---|---| separator carries no data
        row = _TABLE_ROW.match(line)
        if row:
            flush_paragraph()
            table.append([_unescape(cell.strip()) for cell in row.group(1).split("|")])
            continue
        flush_table()

        heading = _HEADING.match(line)
        if heading:
            flush_paragraph()
            blocks.append(
                Block("heading", _unescape(heading.group(2).strip()), len(heading.group(1)))
            )
            continue
        bullet = _BULLET.match(line)
        if bullet:
            flush_paragraph()
            blocks.append(Block("bullet", _unescape(bullet.group(1).strip())))
            continue
        if not line.strip():
            flush_paragraph()
            continue
        paragraph.append(line.strip())

    flush_paragraph()
    flush_table()
    return blocks


def missing_writer(suffix: str) -> str | None:
    """Why this format cannot be written here, or None. Checked before any work."""
    from backend.documents.extract import _importable

    ext = suffix.lower()
    if ext in LEGACY:
        target, described = LEGACY[ext]
        return (
            f"{ext} is {described}, which cannot be written directly. Write a .{target} and"
            f" convert it with convert_file(source=..., target_format='{ext.lstrip('.')}')."
        )
    if ext in {".docx", ".pdf"} and not _importable("docx"):
        # .pdf is written by way of a .docx, so it needs the same library.
        return (
            "writing .docx needs python-docx, which is not installed here. It covers Word"
            " documents; install it with 'pip install \"amethyst[documents]\"' and try again."
        )
    if ext == ".xlsx" and not _importable("openpyxl"):
        return (
            "writing .xlsx needs openpyxl, which is not installed here. It covers Excel"
            " workbooks; install it with 'pip install \"amethyst[documents]\"' and try again."
        )
    if ext == ".pptx" and not _importable("pptx"):
        # Unlike reading, there is no honest fallback: a .pptx is a zip of XML
        # that a dozen cross-referenced parts have to agree about, and hand-
        # writing one produces a file PowerPoint refuses to open.
        return (
            "writing .pptx needs python-pptx, which is not installed here. It covers"
            " PowerPoint decks; install it with 'pip install \"amethyst[documents]\"' and try"
            " again. Reading a deck works without it."
        )
    if ext not in WRITABLE:
        return (
            f"{ext or 'that'} is not a document format this writes. It writes .docx, .xlsx,"
            " .pptx and .pdf; plain text and markdown go through write_file."
        )
    return None


def _save_atomically(save, path: Path) -> None:
    """Write via a sibling temp file, then replace.

    `write_file` is not atomic and does not need to be -- a half-written note is
    readable and recoverable. A half-written 5MB workbook is neither, and the
    window is seconds rather than microseconds because saving a document
    re-serialises the whole thing.
    """
    temporary = path.with_name(f".{path.name}.amethyst-tmp")
    try:
        save(str(temporary))
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink(missing_ok=True)


# --------------------------------------------------------------------------
# Creating
# --------------------------------------------------------------------------


def _table_into_docx(document, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(width):
            table.cell(r, c).text = row[c] if c < len(row) else ""


def _create_docx(path: Path, blocks: list[Block]) -> None:
    from docx import Document

    document = Document()
    for block in blocks:
        if block.kind == "heading":
            # level 0 is Word's "Title" style, which is what `extract._heading_level`
            # reads back as markdown level 1. The two mappings are inverses on
            # purpose: a document written from markdown extracts to the same
            # markdown.
            document.add_heading(block.text, level=min(block.level - 1, 9))
        elif block.kind == "bullet":
            document.add_paragraph(block.text, style="List Bullet")
        elif block.kind == "table" and block.rows:
            _table_into_docx(document, block.rows)
        elif block.text:
            document.add_paragraph(block.text)
    _save_atomically(document.save, path)


def _create_xlsx(path: Path, blocks: list[Block]) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    used = False
    row_cursor = 1

    for block in blocks:
        if block.kind == "heading":
            name = re.sub(r"^Sheet:\s*", "", block.text)[:31] or "Sheet"
            if used:
                sheet = workbook.create_sheet(name)
            else:
                sheet.title, used = name, True
            row_cursor = 1
            continue
        rows = block.rows if block.kind == "table" else [[block.text]] if block.text else []
        for row in rows:
            for column, value in enumerate(row, 1):
                if value != "":
                    sheet.cell(row=row_cursor, column=column, value=_coerce(value))
            row_cursor += 1
        used = used or bool(rows)
    _save_atomically(workbook.save, path)


def _coerce(value: str):
    """Numbers as numbers, so a spreadsheet can add up the column it was given."""
    text = value.strip()
    try:
        return int(text)
    except ValueError:
        pass
    try:
        return float(text)
    except ValueError:
        return value


def _create_pptx(path: Path, blocks: list[Block]) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    slide = None
    body = None

    def new_slide(title: str) -> None:
        nonlocal slide, body
        slide = presentation.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8.8), Inches(1.0))
        box.text_frame.text = title
        body = slide.shapes.add_textbox(
            Inches(0.6), Inches(1.6), Inches(8.8), Inches(5.0)
        ).text_frame
        body.text = ""

    notes_mode = False
    for block in blocks:
        if block.kind == "heading" and block.level <= 2:
            notes_mode = False
            new_slide(block.text)
            continue
        if block.kind == "heading" and block.text.lower() == "notes":
            notes_mode = True
            continue
        if slide is None:
            new_slide(block.text or "Slide")
            continue
        text = block.text if block.kind != "table" else "\n".join(" | ".join(r) for r in block.rows)
        if not text:
            continue
        if notes_mode:
            frame = slide.notes_slide.notes_text_frame
            frame.text = f"{frame.text}\n{text}".strip()
        else:
            line = f"- {text}" if block.kind == "bullet" else text
            body.text = f"{body.text}\n{line}".strip() if body.text else line
    if slide is None:
        new_slide("Slide")
    _save_atomically(presentation.save, path)


def create(path: Path, markdown: str) -> str:
    """Write one document from markdown. Returns a sentence describing what happened."""
    refusal = missing_writer(path.suffix)
    if refusal:
        raise DocumentError(refusal)

    blocks = parse_markdown(markdown)
    if not blocks:
        raise DocumentError("there is nothing in `content` to write")

    path.parent.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        _create_docx(path, blocks)
        return f"wrote {path}"
    if suffix == ".xlsx":
        _create_xlsx(path, blocks)
        return f"wrote {path}"
    if suffix == ".pptx":
        _create_pptx(path, blocks)
        return f"wrote {path}"
    raise DocumentError(f"{suffix} is not written here")  # .pdf is handled by the tool


# --------------------------------------------------------------------------
# Editing
# --------------------------------------------------------------------------


def _workbook_for_writing(path: Path):
    from openpyxl import load_workbook

    # data_only=False, and this is the whole reason the line exists. Extraction
    # opens workbooks with data_only=True to get the cached *value* of a
    # formula, which is what a reader wants. Saving a handle opened that way
    # writes those cached values back over the formulas -- every formula in the
    # file, replaced by whatever it last evaluated to, silently. A fresh
    # workbook is loaded here rather than reusing anything the reader opened.
    return load_workbook(str(path), data_only=False)


def set_cell(path: Path, sheet: str | None, cell: str, value) -> str:
    workbook = _workbook_for_writing(path)
    target = _sheet(workbook, sheet)
    before = target[cell].value
    target[cell] = _coerce(value) if isinstance(value, str) else value
    _save_atomically(workbook.save, path)
    was = "empty" if before is None else repr(before)
    return f"set {target.title}!{cell} to {value!r} (was {was})"


def append_row(path: Path, sheet: str | None, values: list) -> str:
    workbook = _workbook_for_writing(path)
    target = _sheet(workbook, sheet)
    target.append([_coerce(v) if isinstance(v, str) else v for v in values])
    _save_atomically(workbook.save, path)
    return f"appended a row to {target.title}, now {target.max_row} rows"


def _sheet(workbook, name: str | None):
    if name is None:
        return workbook.active
    if name not in workbook.sheetnames:
        raise DocumentError(
            f"no sheet named {name!r}. This workbook has: {', '.join(workbook.sheetnames)}"
        )
    return workbook[name]


def _paragraph_text(paragraph) -> str:
    return "".join(run.text for run in paragraph.runs) or paragraph.text


def _set_paragraph_text(paragraph, text: str) -> None:
    """Replace a paragraph's text, flattening the formatting inside it.

    python-docx splits a paragraph into runs at every formatting boundary, so
    "the **quarterly** report" is three runs and a run-by-run replace never
    matches a string spanning the bold word -- it reports success and changes
    nothing. Working at paragraph level always matches; the cost is that the
    edited paragraph takes the formatting of its first run throughout. That cost
    is stated in the tool description and in the result, rather than left to be
    discovered in Word.
    """
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = text


def _docx_paragraphs(document):
    from docx.table import Table

    from backend.documents.extract import _docx_blocks

    for block in _docx_blocks(document):
        if isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
        else:
            yield block


def replace_text(path: Path, old: str, new: str, replace_all: bool) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        paragraphs = list(_docx_paragraphs(document))
        count = sum(_paragraph_text(p).count(old) for p in paragraphs)
        _guard_count(count, old, path, replace_all)
        remaining = count if replace_all else 1
        for paragraph in paragraphs:
            if remaining <= 0:
                break
            text = _paragraph_text(paragraph)
            if old not in text:
                continue
            here = text.count(old) if replace_all else 1
            _set_paragraph_text(paragraph, text.replace(old, new, here))
            remaining -= here
        _save_atomically(document.save, path)
        return (
            f"replaced {count if replace_all else 1} of {count} in {path.name}."
            " Formatting inside an edited paragraph is flattened to its first run."
        )

    if suffix == ".pptx":
        refusal = missing_writer(suffix)
        if refusal:
            raise DocumentError(refusal)
        from pptx import Presentation

        presentation = Presentation(str(path))
        frames = [
            shape.text_frame
            for slide in presentation.slides
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
        ]
        count = sum(frame.text.count(old) for frame in frames)
        _guard_count(count, old, path, replace_all)
        remaining = count if replace_all else 1
        for frame in frames:
            if remaining <= 0:
                break
            for paragraph in frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs)
                if old not in text or remaining <= 0:
                    continue
                here = text.count(old) if replace_all else 1
                _set_paragraph_text(paragraph, text.replace(old, new, here))
                remaining -= here
        _save_atomically(presentation.save, path)
        return f"replaced {count if replace_all else 1} of {count} in {path.name}"

    raise DocumentError(f"replace_text does not work on {suffix or 'that'}")


def _guard_count(count: int, old: str, path: Path, replace_all: bool) -> None:
    """`edit_file`'s contract, word for word, because it is the one already learned."""
    if count == 0:
        raise DocumentError(f"old_string not found in {path}")
    if count > 1 and not replace_all:
        raise DocumentError(
            f"old_string appears {count} times in {path}; pass replace_all or use a longer,"
            " unique string"
        )


def append_paragraph(path: Path, text: str, style: str | None) -> str:
    from docx import Document

    document = Document(str(path))
    document.add_paragraph(text, style=style or None)
    _save_atomically(document.save, path)
    return f"appended a paragraph to {path.name}"
