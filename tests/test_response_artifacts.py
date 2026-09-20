from __future__ import annotations

import io
from fastapi.testclient import TestClient
import docx

from backend.api.main import app
from backend.db.repositories import (
    ConversationRepository,
    MessageRepository,
    ResponseArtifactRepository,
)
from backend.web.exporter import markdown_to_docx, markdown_to_html


def test_artifact_repository_lifecycle():
    conv_repo = ConversationRepository()
    msg_repo = MessageRepository()
    artifact_repo = ResponseArtifactRepository()

    # 1. Setup conversation & assistant message
    cid = conv_repo.create("test-model", "Test Conversation")
    mid = msg_repo.append(cid, "assistant", "# Initial Response\n\nHere is the initial text.")

    # 2. Get or create artifact
    artifact = artifact_repo.get_or_create(cid, mid, "# Initial Response\n\nHere is the initial text.")
    assert artifact is not None
    assert artifact["version"] == 1
    assert artifact["original_content"] == "# Initial Response\n\nHere is the initial text."
    assert artifact["current_content"] == "# Initial Response\n\nHere is the initial text."
    assert len(artifact["versions"]) == 1
    assert artifact["versions"][0]["author"] == "assistant"

    # 3. Save a user edit (version 2)
    v2_content = "# Initial Response\n\nHere is the edited text by user."
    updated = artifact_repo.save_version(artifact["id"], v2_content, author="user", change_summary="Fixed typo")
    assert updated["version"] == 2
    assert updated["current_content"] == v2_content
    assert len(updated["versions"]) == 2
    assert updated["versions"][0]["version"] == 2
    assert updated["versions"][0]["change_summary"] == "Fixed typo"

    # Check that messages table was updated synchronously
    msg_rows = msg_repo.history(cid)
    assert len(msg_rows) == 1
    assert msg_rows[0].content == v2_content

    # 4. Save another edit (version 3)
    v3_content = "# Initial Response\n\nHere is version 3 with more information."
    artifact_repo.save_version(artifact["id"], v3_content, author="user", change_summary="Added details")
    curr = artifact_repo.get(artifact["id"])
    assert curr["version"] == 3
    assert len(curr["versions"]) == 3

    # 5. Revert back to version 1
    reverted = artifact_repo.revert_to_version(artifact["id"], 1)
    assert reverted["version"] == 4
    assert reverted["current_content"] == "# Initial Response\n\nHere is the initial text."
    assert "Reverted to version 1" in reverted["versions"][0]["change_summary"]

    # Verify messages table has the reverted content
    msg_rows = msg_repo.history(cid)
    assert msg_rows[0].content == "# Initial Response\n\nHere is the initial text."


def test_markdown_to_docx_preserves_structure():
    md = """# Document Title

This is an introduction paragraph with **bold text** and *italic text*.

## Section 1: Features
- First bullet item
- Second bullet item
- Third bullet item

### Ordered Steps
1. Step one
2. Step two

### Code Example
```python
def greet(name):
    return f"Hello {name}"
```

### Comparative Table
| Feature | Amethyst | Traditional |
| --- | --- | --- |
| Artifacts | Supported | None |
| Versioning | Automatic | Manual |
"""

    docx_bytes = markdown_to_docx(md, title="Test Document")
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 1000

    # Parse with python-docx to verify content
    doc = docx.Document(io.BytesIO(docx_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    assert "Document Title" in paragraphs
    assert any("introduction paragraph" in p for p in paragraphs)
    assert any("First bullet item" in p for p in paragraphs)
    assert any("Step one" in p for p in paragraphs)

    # Check table
    assert len(doc.tables) >= 1
    assert any("Feature" in [cell.text for cell in row.cells] for tbl in doc.tables for row in tbl.rows)
    assert any("Amethyst" in [cell.text for cell in row.cells] for tbl in doc.tables for row in tbl.rows)


def test_markdown_to_html():
    md = """# My Title
Paragraph text with `inline code`.

- Item A
- Item B

| Col1 | Col2 |
| --- | --- |
| ValA | ValB |
"""
    html = markdown_to_html(md, title="Custom Title")
    assert "<!DOCTYPE html>" in html
    assert "<title>Custom Title</title>" in html
    assert "<h1>My Title</h1>" in html
    assert "<table>" in html
    assert "<th>Col1</th>" in html
    assert "<td>ValA</td>" in html
    assert "<li>Item A</li>" in html


def test_artifact_api_endpoints():
    client = TestClient(app)
    conv_repo = ConversationRepository()
    msg_repo = MessageRepository()

    cid = conv_repo.create("test-model", "API Test Conversation")
    mid = msg_repo.append(cid, "assistant", "### Original Assistant Text\nDetailed analysis here.")

    # 1. GET artifact
    res = client.get(f"/api/conversations/{cid}/messages/{mid}/artifact")
    assert res.status_code == 200
    data = res.json()
    assert data["version"] == 1
    assert "Original Assistant Text" in data["current_content"]

    # 2. POST artifact update
    new_text = "### Original Assistant Text\nUpdated with corrections."
    res = client.post(
        f"/api/conversations/{cid}/messages/{mid}/artifact",
        json={"content": new_text, "change_summary": "Fixed details"},
    )
    assert res.status_code == 200
    updated_data = res.json()
    assert updated_data["version"] == 2
    assert updated_data["current_content"] == new_text

    # 3. POST revert
    res = client.post(
        f"/api/conversations/{cid}/messages/{mid}/artifact/revert",
        json={"version": 1},
    )
    assert res.status_code == 200
    rev_data = res.json()
    assert rev_data["version"] == 3
    assert "Detailed analysis here." in rev_data["current_content"]

    # 4. POST export/docx
    res = client.post(
        "/api/export/docx",
        json={"markdown": "# Export Test\nSample text", "title": "Docx Test"},
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert len(res.content) > 1000

    # 5. POST ai/transform - shorten
    res = client.post(
        "/api/ai/transform",
        json={
            "text": "This is a rather long and needlessly convoluted sentence that contains extra words.",
            "action": "shorten",
        },
    )
    assert res.status_code == 200
    transform_data = res.json()
    assert "transformed" in transform_data
    assert len(transform_data["transformed"]) > 0

    # 6. POST ai/transform - rewrite
    res = client.post(
        "/api/ai/transform",
        json={"text": "Original raw draft with basic thoughts.", "action": "rewrite"},
    )
    assert res.status_code == 200
    assert "transformed" in res.json()
    assert res.json()["transformed"] != ""

    # 7. POST ai/transform - custom instruction (bullet list)
    res = client.post(
        "/api/ai/transform",
        json={
            "text": "Feature A is great. Feature B is faster. Feature C is secure.",
            "instruction": "Convert to bullet points",
            "action": "rewrite",
        },
    )
    assert res.status_code == 200
    res_text = res.json()["transformed"]
    assert "- " in res_text or "Feature" in res_text

    # 8. POST ai/transform - expand
    res = client.post(
        "/api/ai/transform",
        json={"text": "System operational status.", "action": "expand"},
    )
    assert res.status_code == 200
    assert len(res.json()["transformed"]) > len("System operational status.")

